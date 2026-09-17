import json
import sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

REQUIRED = ["세금계산서날짜", "수량", "단가(원)", "금액(원)", "납품업체명", "모델명", "Sub Category 1", "Sub Category 3", "총C/D금액", "C/D%"]

def serial(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value

def has_value(value):
    return value is not None and str(value).strip() not in ("", "-")

def positive_number(value):
    if not has_value(value):
        return False
    try:
        return float(str(value).replace(",", "")) > 0
    except (TypeError, ValueError):
        return False

def clean(value):
    return value if has_value(value) else None

def main():
    if len(sys.argv) != 2:
        raise ValueError("입력 파일이 없습니다.")
    path = Path(sys.argv[1])
    wb = load_workbook(path, data_only=True, read_only=True)
    if "CD집계표(Total)" not in wb.sheetnames:
        raise ValueError("CD집계표(Total) 시트를 찾을 수 없습니다.")
    ws = wb["CD집계표(Total)"]
    all_rows = ws.iter_rows(values_only=True)
    for _ in range(3):
        next(all_rows)
    headers = [str(value).strip() if value is not None else "" for value in next(all_rows)]
    missing = [name for name in REQUIRED if name not in headers]
    if not any(name in headers for name in ("PO번호", "PO 번호")):
        missing.append("PO번호 또는 PO 번호")
    if missing:
        print(json.dumps({"ok": False, "errorType": "missing_column", "missing": missing}, ensure_ascii=True))
        return
    index = {name: i for i, name in enumerate(headers)}
    rows = []
    skipped_template_rows = 0
    for row_number, raw_values in enumerate(all_rows, start=5):
        values = [serial(value) for value in raw_values]
        if not any(value not in (None, "") for value in values):
            continue
        get = lambda name: values[index[name]] if index[name] < len(values) else None
        invoice_date = get("세금계산서날짜")
        purchase_amount = get("금액(원)")
        # 완료 구매 이력은 세금계산서 날짜 또는 실제 금액이 있어야 합니다.
        # 서식·수식만 남아 있는 템플릿 행은 분석 대상에서 제외합니다.
        if not has_value(invoice_date) and not positive_number(purchase_amount):
            skipped_template_rows += 1
            continue
        rows.append({
            "sourceRow": row_number,
            # CD 집계표의 원본 구매 PO번호를 우선 사용한다. AP 관리용 "PO 번호"는
            # 원본 PO가 비어 있는 경우에만 보조 값으로 사용한다.
            "poNumber": clean(get("PO번호")) if "PO번호" in index and clean(get("PO번호")) else (clean(get("PO 번호")) if "PO 번호" in index else None),
            "taxInvoiceDate": clean(invoice_date),
            "quantity": clean(get("수량")),
            "unitPrice": clean(get("단가(원)")),
            "purchaseAmount": clean(purchase_amount),
            "supplierName": clean(get("납품업체명")),
            "modelName": clean(get("모델명")),
            "categoryLarge": clean(get("Sub Category 1")),
            "categorySmall": clean(get("Sub Category 3")),
            "costDownAmount": clean(get("총C/D금액")),
            "costDownRate": clean(get("C/D%")),
            "buyer": clean(get("Buyer")),
            "purchasePurpose": clean(get("구매용도")),
            "requester": clean(get("요청자")),
        })
    warningRows = sum(1 for row in rows if not row["poNumber"])
    print(json.dumps({"ok": True, "sheet": "CD집계표(Total)", "headers": headers, "rows": rows, "warningRows": warningRows, "skippedTemplateRows": skipped_template_rows}, ensure_ascii=True))

if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"ok": False, "errorType": "parse_error", "error": str(exc)}, ensure_ascii=True))
