import json
import sys
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook


def serial(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def text(value):
    return "" if value is None else str(value).strip()


def sheet_rows(ws, header_row):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < header_row:
        return [], []
    headers = [text(v) for v in rows[header_row - 1]]
    return headers, rows[header_row:]


def getter(headers, row):
    index = {h: i for i, h in enumerate(headers) if h}

    def get(*names):
        for name in names:
            if name in index and index[name] < len(row):
                return serial(row[index[name]])
        return None

    return get


def parse_contracts(wb):
    if "Ariba입력대상" not in wb.sheetnames:
        return [], ["Ariba입력대상 시트"]
    ws = wb["Ariba입력대상"]
    headers, rows = sheet_rows(ws, 2)
    needed = ["계약ID", "업체명", "전체 설명(Description)", "지불월도", "총합계 (VAT포함)", "담당자"]
    missing = [name for name in needed if name not in headers]
    if missing:
        return [], missing
    result = []
    for source_row, raw in enumerate(rows, start=3):
        if not any(v not in (None, "") for v in raw):
            continue
        get = getter(headers, raw)
        contract_id = text(get("계약ID"))
        supplier = text(get("업체명"))
        description = text(get("전체 설명(Description)"))
        if not contract_id and not supplier and not description:
            continue
        result.append({
            "sourceRow": source_row,
            "isBrightRow": ws.cell(source_row, 1).fill.fill_type is None,
            "poNumber": text(get("PO 번호", "PO NO")),
            "contractId": contract_id,
            "supplierName": supplier or "미입력",
            "product": description or "미입력",
            "type": text(get("Compliance Contract Type", "계약 유형")) or "유지보수",
            "owner": text(get("담당자")),
            "customerName": text(get("고객명")),
            "startDate": get("고객계약시작일", "계약 시작일"),
            "endDate": get("고객계약만료", "고객계약만료일", "계약 만료일"),
            "supplierEndDate": get("Ariba 계약만료일"),
            "billingCycle": text(get("지불월도")) or "미입력",
            "currency": "KRW",
            "amountInclVat": get("총합계 (VAT포함)"),
            "amountExclVat": get("수락됨\n- Amount(VAT별도)", "Amount(VAT별도)"),
            "salesType": text(get("판매형태")),
            "account": text(get("ACCOUNT")),
            "paymentExcluded": text(get("지불여부")) in ("대상외", "지급 제외"),
            "maintenanceEligible": True,
        })
    return result, []


def parse_payments(wb):
    # 지급 대상은 유지보수지불건 시트를 사용하지 않고 Ariba입력대상에서 산정한다.
    return [], []
    if "유지보수지불건" not in wb.sheetnames:
        return [], []
    headers, rows = sheet_rows(wb["유지보수지불건"], 4)
    result = []
    for source_row, raw in enumerate(rows, start=5):
        if not any(v not in (None, "") for v in raw):
            continue
        get = getter(headers, raw)
        supplier = text(get("업체명"))
        description = text(get("Description"))
        if not supplier and not description:
            continue
        result.append({
            "sourceRow": source_row,
            "contractId": text(get("계약ID")),
            "supplierName": supplier,
            "product": description,
            "paymentMonth": text(get("지불월도")),
            "amountExclVat": get("Amount"),
            "amountInclVat": get("Invoice Amount", "Amount"),
            "receiptStatus": text(get("인수증 상태", "입력여부")),
            "invoiceNumber": text(get("Invoice Num", "인보이스 번호")),
            "apSlipNumber": text(get("AP 전표번호")),
            "paymentStatus": "AP 전표 완료" if text(get("입력여부")) == "완료" else "인수증 발행",
            "account": text(get("ACCOUNT")),
            "team": text(get("팀명")),
            "owner": text(get("담당자")),
        })
    return result, []


def parse_shipments(wb):
    if "쉽컴펌" not in wb.sheetnames:
        return [], []
    headers, rows = sheet_rows(wb["쉽컴펌"], 3)
    result = []
    for source_row, raw in enumerate(rows, start=4):
        if not any(v not in (None, "") for v in raw):
            continue
        get = getter(headers, raw)
        po = text(get("PO NO"))
        code = text(get("CODE"))
        if not po and not code:
            continue
        status = text(get("Status"))
        normalized = "쉽컴펌 완료" if "완료" in status else (status or "출고대기")
        result.append({
            "sourceRow": source_row,
            "poNumber": po,
            "orderNumber": text(get("주문번호")),
            "deliveryPlace": text(get("납품처(고객)")),
            "itemCode": code,
            "quantity": get("수량"),
            "unitPrice": get("단가"),
            "amount": get("금액"),
            "serialNumber": text(get("시리얼넘버")),
            "receivedDate": get("PO입고일"),
            "shipmentStatus": normalized,
            "assignee": text(get("담당자(지역)")),
            "memo": text(get("비고")),
        })
    return result, []


def main():
    if len(sys.argv) != 2:
        raise ValueError("입력 파일이 없습니다.")
    # 행 음영(밝은 행/음영 행)을 판별해야 하므로 read_only를 사용하지 않는다.
    wb = load_workbook(Path(sys.argv[1]), data_only=True, read_only=False, keep_vba=True)
    contracts, c_missing = parse_contracts(wb)
    payments, p_missing = parse_payments(wb)
    shipments, s_missing = parse_shipments(wb)
    missing = sorted(set(c_missing + p_missing + s_missing))
    if missing:
        print(json.dumps({"ok": False, "errorType": "missing_column", "missing": missing}, ensure_ascii=True))
        return
    warnings = []
    for row in contracts:
        if not row["contractId"]:
            warnings.append({"sheet": "Ariba입력대상", "sourceRow": row["sourceRow"], "reason": "계약ID 누락"})
    for row in shipments:
        if not row["poNumber"]:
            warnings.append({"sheet": "쉽컴펌", "sourceRow": row["sourceRow"], "reason": "PO 번호 누락"})
    print(json.dumps({
        "ok": True,
        "sheets": [name for name in ["Ariba입력대상", "유지보수지불건", "쉽컴펌"] if name in wb.sheetnames],
        "contracts": contracts,
        "payments": payments,
        "shipments": shipments,
        "warnings": warnings,
        "summary": {"contracts": len(contracts), "payments": len(payments), "shipments": len(shipments), "warnings": len(warnings)},
    }, ensure_ascii=True, default=str))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"ok": False, "errorType": "parse_error", "error": str(exc)}, ensure_ascii=True))
